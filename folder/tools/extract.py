"""Extract every slide/page of the course files into work/<doc>/source.json plus a JPEG per slide.

Usage: python3 tools/extract.py <file> [<file> ...]   (run from folder/)
"""
import json
import re
import subprocess
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

ROOT = Path(__file__).resolve().parent.parent
WORK = ROOT / "work"
FOOTER = re.compile(r"^©\s*20\d\d Amazon Web Services.*$|^\d{1,3}$")


def clean(lines):
    out = []
    for ln in lines:
        ln = ln.replace("\u000b", " ").strip()
        if ln and not FOOTER.match(ln):
            out.append(ln)
    return out


def shape_texts(shape):
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for s in sorted(shape.shapes, key=lambda s: ((s.top or 0), (s.left or 0))):
            yield from shape_texts(s)
        return
    if shape.has_text_frame:
        for p in shape.text_frame.paragraphs:
            t = "".join(r.text for r in p.runs).strip()
            if t:
                yield ("  " * p.level) + t
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            yield " | ".join(c.text.strip().replace("\n", " ") for c in row.cells)


def render(pdf: Path, out: Path):
    out.mkdir(parents=True, exist_ok=True)
    if not any(out.glob("*.jpg")):
        subprocess.run(["pdftoppm", "-r", "80", "-jpeg", "-jpegopt", "quality=70", str(pdf), str(out / "p")], check=True)


def from_pptx(path: Path):
    prs = Presentation(str(path))
    units = []
    for i, slide in enumerate(prs.slides, 1):
        shapes = sorted(slide.shapes, key=lambda s: ((s.top or 0), (s.left or 0)))
        texts = clean([t for s in shapes for t in shape_texts(s)])
        notes = []
        if slide.has_notes_slide:
            notes = clean(slide.notes_slide.notes_text_frame.text.split("\n"))
        hidden = slide._element.get("show") == "0"
        units.append({"id": f"S{i:03d}", "texts": texts, "notes": notes, "hidden": hidden})
    pdf = WORK / "pptxpdf" / (path.stem + ".pdf")
    # SmartArt / diagram text is invisible to python-pptx, so also keep the text of the rendered slide
    for i, u in enumerate(units, 1):
        txt = subprocess.run(["pdftotext", "-f", str(i), "-l", str(i), str(pdf), "-"], capture_output=True, text=True).stdout
        u["rendered"] = clean(txt.split("\n"))
    return units, pdf


def from_pdf(path: Path):
    n = int(re.search(r"Pages:\s+(\d+)", subprocess.run(["pdfinfo", str(path)], capture_output=True, text=True).stdout).group(1))
    units = []
    for i in range(1, n + 1):
        txt = subprocess.run(["pdftotext", "-f", str(i), "-l", str(i), str(path), "-"], capture_output=True, text=True).stdout
        units.append({"id": f"P{i:03d}", "texts": clean(txt.split("\n")), "notes": [], "hidden": False, "rendered": []})
    return units, path


def main(files):
    for f in files:
        path = (ROOT / f).resolve()
        units, pdf = from_pptx(path) if path.suffix == ".pptx" else from_pdf(path)
        d = WORK / path.stem
        d.mkdir(parents=True, exist_ok=True)
        render(pdf, d / "img")
        imgs = sorted((d / "img").glob("*.jpg"))
        assert len(imgs) == len(units), f"{f}: {len(imgs)} images vs {len(units)} units"
        for u, img in zip(units, imgs):
            u["img"] = f"img/{img.name}"
        (d / "source.json").write_text(json.dumps({"file": path.name, "units": units}, ensure_ascii=False, indent=1))
        words = sum(len(" ".join(u["texts"] + u["notes"]).split()) for u in units)
        print(f"{path.name}: {len(units)} units, {words} words, hidden={sum(u['hidden'] for u in units)}")


if __name__ == "__main__":
    main(sys.argv[1:])
